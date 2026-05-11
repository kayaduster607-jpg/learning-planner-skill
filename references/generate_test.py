# -*- coding: utf-8 -*-
"""
学习规划师 — 水平测评卷 Word 文档生成器
=========================================
通用模板脚本，修改 main() 中的 questions 列表即可适配任意领域。

题型：
  - "choice" : 选择题（需 options 列表）
  - "judge"  : 判断题
  - "fill"   : 填空题（留一行横线）
  - "short"  : 简答题（留三行横线）
  - "write"  : 写作/翻译题（留六行横线，自动归入选做部分）

写作/翻译类自动归入"选做部分"，其余归入"必做部分"。
预计用时自动计算。
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os

# ===== 配置 =====
FONT_CN = "Microsoft YaHei"
FONT_EN = "Calibri"
PAGE_MARGIN = Cm(2.5)

# 每种题型预计用时（分钟）
TIME_PER_TYPE = {
    "choice": 1,
    "judge": 1,
    "fill": 1.5,
    "short": 2.5,
    "write": 12,
}

OPTIONAL_TYPES = {"write"}


def set_run_font(run, size=11, bold=False, color=None, italic=False):
    """设置 run 字体，确保中文不乱码"""
    run.font.name = FONT_EN
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color
    rPr = run._element.find(qn("w:rPr"))
    if rPr is None:
        rPr = run._element.makeelement(qn("w:rPr"), {})
        run._element.insert(0, rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), FONT_CN)


def add_text(doc, text, size=11, bold=False, color=None, align=None, italic=False, space_after=None):
    """添加段落"""
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return p


def add_blank_lines(doc, count=1):
    """添加空白横线（答题区）"""
    for _ in range(count):
        p = doc.add_paragraph()
        run = p.add_run("_" * 75)
        set_run_font(run, size=10, color=RGBColor(0xBB, 0xBB, 0xBB))


def add_question(doc, q):
    """添加一道题"""
    num = q["num"]
    text = q["text"]
    qtype = q["type"]
    note = q.get("note", "")
    options = q.get("options", [])

    # 题号 + 题干
    p = doc.add_paragraph()
    run = p.add_run(f"{num}. {text}")
    set_run_font(run, size=11, bold=True)

    # 选项
    if options:
        for opt in options:
            op = doc.add_paragraph()
            op.paragraph_format.left_indent = Cm(1.2)
            run = op.add_run(opt)
            set_run_font(run, size=11)

    # 答题区
    if qtype == "fill":
        add_blank_lines(doc, 1)
    elif qtype == "short":
        add_blank_lines(doc, 3)
    elif qtype == "write":
        add_blank_lines(doc, 6)

    # 考察标签
    if note:
        tag = doc.add_paragraph()
        tag.paragraph_format.space_after = Pt(14)
        run = tag.add_run(note)
        set_run_font(run, size=9, color=RGBColor(0x88, 0x88, 0x88), italic=True)


def calc_time(questions):
    """计算必做/选做预计用时"""
    req = [q for q in questions if q["type"] not in OPTIONAL_TYPES]
    opt = [q for q in questions if q["type"] in OPTIONAL_TYPES]
    req_time = sum(TIME_PER_TYPE.get(q["type"], 1) for q in req)
    opt_time = sum(TIME_PER_TYPE.get(q["type"], 1) for q in opt)
    return round(req_time), round(opt_time)


def generate_test(title, domain, questions, output_dir=None):
    """
    生成测评卷 Word 文档

    Args:
        title:    文档标题（如 "Python 编程水平测评卷"）
        domain:   领域名称（如 "Python 编程"），用于文件名
        questions: 题目列表，每题为 dict：
            {
                "num": 1,
                "type": "choice" | "judge" | "fill" | "short" | "write",
                "text": "题目内容",
                "options": ["A. ...", "B. ..."],  # 选择题必填
                "note": "[考察：XXX]"              # 可选
            }
        output_dir: 输出目录，默认为脚本所在目录

    Returns:
        str: 生成的文件绝对路径
    """
    doc = Document()

    # 设置默认字体
    style = doc.styles["Normal"]
    style.font.name = FONT_EN
    style.font.size = Pt(11)
    rPr = style.element.find(qn("w:rPr"))
    if rPr is None:
        rPr = style.element.makeelement(qn("w:rPr"), {})
        style.element.append(rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), FONT_CN)

    # 页面边距
    for section in doc.sections:
        section.top_margin = PAGE_MARGIN
        section.bottom_margin = PAGE_MARGIN
        section.left_margin = PAGE_MARGIN
        section.right_margin = PAGE_MARGIN

    # 分区
    required = [q for q in questions if q["type"] not in OPTIONAL_TYPES]
    optional = [q for q in questions if q["type"] in OPTIONAL_TYPES]
    req_time, opt_time = calc_time(questions)

    # ===== 标题 =====
    title_p = doc.add_heading("", level=0)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(title)
    set_run_font(run, size=22, bold=True)

    add_text(doc, "请诚实作答，不确定的题目可以写\"不清楚\"。",
             size=11, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x55, 0x55, 0x55))
    add_text(doc, "测评目的是精准定位你的当前水平，不是考试，答错很正常。",
             size=11, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x55, 0x55, 0x55))
    add_text(doc, "姓名：______________    日期：______________",
             size=11, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=20)

    # ===== 必做部分 =====
    header = doc.add_heading(f"必做部分（约 {req_time} 分钟 · {len(required)} 题）", level=1)
    for run in header.runs:
        set_run_font(run, size=16, bold=True)
    add_text(doc, "请优先完成本部分，完成后即可提交获得基础水平定级。",
             size=10, color=RGBColor(0x66, 0x66, 0x66), space_after=12)

    for q in required:
        add_question(doc, q)

    # ===== 选做部分 =====
    if optional:
        doc.add_page_break()
        header = doc.add_heading(f"选做部分（如时间充裕，约 {opt_time} 分钟 · {len(optional)} 题）", level=1)
        for run in header.runs:
            set_run_font(run, size=16, bold=True)
        add_text(doc, "以下题目用于更精准评估输出能力，不做强制要求。如有时间建议作答，会进一步细化学习计划。",
                 size=10, color=RGBColor(0x66, 0x66, 0x66), space_after=12)
        for q in optional:
            add_question(doc, q)

    # ===== 底部 =====
    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("完成后提交，将为你分析水平并制定学习计划。")
    set_run_font(run, size=11, bold=True, color=RGBColor(0x33, 0x33, 0x33))

    # 保存
    if output_dir is None:
        output_dir = os.path.dirname(os.path.abspath(__file__))
    os.makedirs(output_dir, exist_ok=True)
    filename = f"{domain}水平测评卷.docx"
    output_path = os.path.join(output_dir, filename)
    doc.save(output_path)
    return output_path


# ===== 示例：英语六级 =====
def main():
    title = "英语六级水平测评卷"
    domain = "英语六级"

    questions = [
        # --- 必做部分 ---
        {
            "num": 1, "type": "choice",
            "text": "The new policy has had a ______ effect on the local economy, bringing in more investment and creating jobs.",
            "options": ["A. detrimental", "B. negligible", "C. profound", "D. adverse"],
            "note": "[考察：六级高频形容词辨析]"
        },
        {
            "num": 2, "type": "fill",
            "text": "请写出以下单词的中文含义（能写多少写多少）：\n    deteriorate / contemplate / indispensable / vicinity / spontaneous",
            "note": "[考察：六级核心词汇量摸底]"
        },
        {
            "num": 3, "type": "judge",
            "text": "判断对错：在英语中，\"suggest\"后面的宾语从句必须用虚拟语气（should + 动词原形），should 可以省略。",
            "note": "[考察：虚拟语气基础规则]"
        },
        {
            "num": 4, "type": "fill",
            "text": "完成句子：If I ______ (know) the answer yesterday, I would have told you.",
            "note": "[考察：虚拟语气——与过去事实相反]"
        },
        {
            "num": 5, "type": "choice",
            "text": "下面哪个句子的语法是正确的？",
            "options": [
                "A. Not only he passed the exam, but also got a scholarship.",
                "B. Not only did he pass the exam, but he also got a scholarship.",
                "C. Not only did he pass the exam, but also got a scholarship."
            ],
            "note": "[考察：倒装句基础]"
        },
        {
            "num": 6, "type": "choice",
            "text": "阅读以下段落，回答问题：\n\nThe rapid advancement of artificial intelligence has sparked intense debate about its impact on employment. While some economists argue that AI will create more jobs than it eliminates, others contend that the displacement of workers in certain sectors could lead to significant social unrest. The truth likely lies somewhere in between — AI will transform the nature of work rather than simply destroying it.\n\n作者对 AI 影响就业的态度是什么？",
            "options": [
                "A. 完全乐观",
                "B. 完全悲观",
                "C. 客观中立，认为影响是复杂的",
                "D. 没有表达态度"
            ],
            "note": "[考察：阅读理解——作者态度判断]"
        },
        {
            "num": 7, "type": "choice",
            "text": "\"The committee decided to ______ the proposal until more data was available.\" 请选择最合适的词：",
            "options": ["A. postpone", "B. abandon", "C. endorse", "D. accelerate"],
            "note": "[考察：六级动词辨析 + 语境理解]"
        },
        {
            "num": 8, "type": "choice",
            "text": "以下哪句话使用了\"倒装\"结构？",
            "options": [
                "A. He seldom goes to the library.",
                "B. Seldom does he go to the library.",
                "C. He does seldom go to the library."
            ],
            "note": "[考察：部分倒装识别]"
        },
        {
            "num": 9, "type": "fill",
            "text": "\"The scientist devoted all her life to ______ (explore) the mysteries of the ocean.\" 括号中应填入什么形式？",
            "note": "[考察：非谓语动词——介词 + doing]"
        },
        {
            "num": 10, "type": "short",
            "text": "请简要解释\"虚拟语气\"的三种基本类型（与现在、过去、将来事实相反），并各举一个例子。",
            "note": "[考察：虚拟语气综合理解]"
        },

        # --- 选做部分 ---
        {
            "num": 11, "type": "write",
            "text": "将以下中文翻译成英文：\n\n\"随着互联网的普及，越来越多的人开始在网上购物，这不仅改变了人们的消费习惯，也对传统零售业产生了深远的影响。\"",
            "note": "[考察：翻译能力——六级翻译典型句式]"
        },
        {
            "num": 12, "type": "write",
            "text": "请用英语写一段 80-100 词的短文，主题：\n\n\"The Importance of Lifelong Learning\"（终身学习的重要性）\n\n要求：至少使用 2 个六级水平的词汇或短语。",
            "note": "[考察：写作能力——组织结构 + 词汇运用 + 语法准确性]"
        },
        {
            "num": 13, "type": "write",
            "text": "请将以下中文翻译成英文：\n\n\"中国的高铁网络是世界上最大的，它极大地缩短了城市之间的旅行时间，促进了区域经济的协调发展。近年来，中国高铁技术也开始走向世界，为其他国家提供了优质的交通解决方案。\"",
            "note": "[考察：翻译能力——长句处理 + 中国特色表达]"
        },
    ]

    # 输出到当前目录，可修改为任意路径
    output_dir = os.path.dirname(os.path.abspath(__file__))
    output_path = generate_test(title, domain, questions, output_dir=output_dir)
    print(f"Word 文档已生成: {output_path}")


if __name__ == "__main__":
    main()
